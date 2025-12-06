import io
import os

from openai import OpenAI
from dotenv import load_dotenv
import chromadb
from chromadb.config import Settings
from pypdf import PdfReader
import docx
from pptx import Presentation

load_dotenv()

_default_model = (
    os.getenv("OPENAI_MODEL")
    or os.getenv("DEEPSEEK_MODEL")
    or "deepseek-chat"
)
_default_embedding_model = (
    os.getenv("OPENAI_EMBEDDING_MODEL")
    or os.getenv("DEEPSEEK_EMBEDDING_MODEL")
    or "text-embedding-3-small"
)


def _get_openai_client():
    """Lazily construct an OpenAI-compatible client.

    Prefers DeepSeek when DEEPSEEK_API_KEY is present, otherwise uses OPENAI_API_KEY.
    Returns None if no supported API key is set so callers can fall back gracefully.
    """

    deepseek_key = os.getenv("DEEPSEEK_API_KEY")
    if deepseek_key:
        base_url = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
        return OpenAI(api_key=deepseek_key, base_url=base_url)

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        return None
    return OpenAI(api_key=api_key)


def _read_txt(content: bytes) -> str:
    try:
        return content.decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _read_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            parts.append(text)
        return "\n".join(parts)
    except Exception:
        return ""


def _read_docx(content: bytes) -> str:
    try:
        f = io.BytesIO(content)
        document = docx.Document(f)
        return "\n".join(p.text for p in document.paragraphs)
    except Exception:
        return ""


def _read_pptx(content: bytes) -> str:
    try:
        prs = Presentation(io.BytesIO(content))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    except Exception:
        return ""


def extract_text_from_uploaded(uploaded):
    extracted = []
    for item in uploaded:
        filename = item.get("filename") or "document"
        content = item.get("content") or b""
        lower = filename.lower()
        if lower.endswith(".pdf"):
            text = _read_pdf(content)
        elif lower.endswith(".docx"):
            text = _read_docx(content)
        elif lower.endswith(".ppt") or lower.endswith(".pptx"):
            text = _read_pptx(content)
        elif lower.endswith(".txt"):
            text = _read_txt(content)
        else:
            text = _read_txt(content)
        if text.strip():
            extracted.append({"filename": filename, "text": text})
    return extracted


def build_rag_collection(docs, topic: str):
    if not docs:
        return None, []

    client = _get_openai_client()
    if client is None:
        # No API key configured – skip building a collection so callers can fall back.
        return None, []

    try:
        chroma = chromadb.Client(Settings(anonymized_telemetry=False))
        collection = chroma.create_collection(name="newsletter-docs")
        ids = []
        texts = []
        metas = []
        for idx, doc in enumerate(docs):
            doc_id = f"doc-{idx}"
            ids.append(doc_id)
            texts.append(doc["text"])
            metas.append({"filename": doc["filename"]})
        embeddings = client.embeddings.create(model=_default_embedding_model, input=texts)
        vectors = [item.embedding for item in embeddings.data]
        collection.add(ids=ids, embeddings=vectors, metadatas=metas, documents=texts)
        return collection, metas
    except Exception:
        # If anything goes wrong with embeddings or Chroma (e.g. quota issues),
        # skip RAG and let callers fall back gracefully.
        return None, []


def retrieve_relevant_context(collection, query: str, top_k: int = 5) -> str:
    if collection is None or not query:
        return ""

    client = _get_openai_client()
    if client is None:
        return ""
    try:
        result = collection.query(
            query_embeddings=[
                client.embeddings.create(model=_default_embedding_model, input=[query]).data[0].embedding
            ],
            n_results=top_k,
        )
        docs = result.get("documents") or []
        flattened = []
        for batch in docs:
            flattened.extend(batch)
        joined = "\n---\n".join(flattened)
        return joined[:8000]
    except Exception:
        # If retrieval or embeddings fail (e.g. due to quota), just return no extra context.
        return ""


def generate_newsletter_with_llm(topic, tone, audience, length, doc_context: str):
    client = _get_openai_client()
    if client is None:
        # Let caller handle the fallback path when no key is configured.
        raise RuntimeError("OPENAI_API_KEY is not configured")
    system = (
        "You are an assistant that creates structured email newsletters in HTML, "
        "based primarily on real content extracted from the user's uploaded documents. "
        "You must extract and summarize the most important news, insights, and takeaways "
        "from the provided document context, then organize them into a clean newsletter. "
        "You must return valid JSON with keys: title, html_body, cta_text, cta_link. "
        "html_body should be a complete responsive email body with inline styles, "
        "including hero/intro, 2-4 sections, a clear summary section, and an explicit CTA."
    )
    user = (
        f"Topic: {topic or 'Smart Newsletter Update'}\n"
        f"Tone: {tone or 'Professional'}\n"
        f"Audience: {audience or 'Customers'}\n"
        f"Length: {length or 'Medium'}\n"
        f"Relevant document context (raw extracted text from PDFs/DOCs/PPT/TXT):\n"
        f"{doc_context or 'None'}\n\n"
        "Use the document context as the primary source of truth. Summarize the key points, "
        "headlines and insights, and relate them to the given topic if provided. "
        "Avoid inventing facts that are not supported by the documents. "
        "Create a newsletter-style email with sections, short headings, and concise summaries, "
        "following the requested tone and length. Return JSON only, no markdown, with keys "
        "exactly: title, html_body, cta_text, cta_link."
    )
    completion = client.chat.completions.create(
        model=_default_model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        response_format={"type": "json_object"},
    )
    content = completion.choices[0].message.content
    return content
